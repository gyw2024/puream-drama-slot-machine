from pathlib import Path
import zipfile, hashlib, json, re
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
src=Path('D:/CodexData/.codex/skills/puream-drama-production-package')
out=Path('D:/Backup/Documents/无限画布/纯梦短剧老虎机/交付文件/资产包制作Skill-20260915')
out.mkdir(parents=True,exist_ok=True)
sections=[
('这是什么', '这是供 Agent 读取执行的技能文件，不是双击运行的安装程序。它用于准备短剧剧本、逐镜视频提示词、参考图片资产及制作清单，最后输出 .pdramapack 文件，供纯梦短剧老虎机的“资产导入”模式使用。ZIP 是技能包，.pdramapack 才是最终制作包，两者不要混用。'),
('怎么使用', '1. 解压 ZIP，保留 puream-drama-production-package 文件夹的完整结构。\n2. 将整个技能文件夹放进所用 Agent 支持的技能目录；已有同名目录时先备份再替换。也可直接在支持本地文件读取的 Agent 中指定解压后的 SKILL.md 路径，让它读取并执行。\n3. 为 Agent 提供商品图片、商品资料，以及剧情要求或完整原稿。\n4. 按下面的示例发出任务，让 Agent 完成制作并交付 .pdramapack。\n5. 在纯梦短剧老虎机中新建“资产导入”项目，通过“导入资产包”选择生成的 .pdramapack，核对内容。\n6. 查看全部提示词并确认后，再继续生成视频。导入文件或 AI 审核完成不等于用户确认生成。'),
('需要准备什么', '• 剧情：原创方向，或者完整的上传／改写原稿；写清必须保留的人物关系、对白与结局。\n• 商品：真实商品原图、名称、卖点、价格、活动与购买入口；没有活动就注明无活动。\n• 工具：能读取本地文件、执行脚本的 Agent；打包脚本需要 Node.js。制作真实图片还需要可用的生图工具和相应额度，技能包本身不附带账号或算力。\n• 输出位置：指定一个保存剧本、图片、提示词和最终制作包的本地文件夹。'),
('可直接复制的任务示例', '请读取并使用我提供的 puream-drama-production-package/SKILL.md，按其中规则制作一个完整短剧资产包。\n创作方式：原创（也可以改成上传剧本或改写）。\n剧情方向：[填写人物关系、主要冲突和结局方向]。\n商品名称：[填写]；商品原图：[附件或文件路径]。\n真实卖点：[填写]；价格：[填写]；活动：[填写，无则写无]；购买入口：[填写]。\n先围绕人物和剧情写作，让商品自然进入剧情，不要写成独立产品广告。\n输出目录：[填写本地路径]。请先给我剧本和全部提示词查看，等我确认后再生成所需图片并打包；不要生成视频。最终交付 .pdramapack、完整剧本、提示词和检查结果，缺少的工具或资料请明确说明。'),
('收到制作包后检查什么', '核对剧情与对白是否完整连贯、人物身份是否一致、场景和物件引用是否正确、商品原图与价格活动是否准确。逐一查看生成的图片，再核对视频提示词的说话人、语气、站位和动作。提示词与图片检查通过，仍不代表未来生成的视频声音和画面一定符合要求；成片需另行查看。'),
('打包内容与注意事项', '技能主文件、references 规则文档、agents 配置及 scripts 辅助脚本均按当前本机版本原样打包。无需手工修改脚本或伪造审核记录。技能内含审核和打包检查逻辑，并非“完全没有校验”的版本。若工具报告缺项，让 Agent 根据具体证据处理并保留已完成内容；不要反复从头生成。\n本次仅交付技能文件及使用说明，未执行真实写作、生图或制作包导入测试。')]
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Cm(1.7)
sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1']:
 st=doc.styles[name];st.font.name='微软雅黑';st.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
doc.styles['Normal'].font.size=Pt(10.5)
doc.add_heading('短剧资产包制作 Skill',0)
doc.add_paragraph('简单使用说明 · 2026年9月15日')
for title,body in sections:
 doc.add_heading(title,1)
 for line in body.split('\n'):doc.add_paragraph(line)
word=out/'资产包制作Skill-简单使用说明.docx';doc.save(word)
readme='# 短剧资产包制作 Skill\n\n'+'\n\n'.join('## '+a+'\n'+b for a,b in sections)
(out/'使用说明.md').write_text(readme,encoding='utf-8')
files=sorted(p for p in src.rglob('*') if p.is_file() and '__pycache__' not in p.parts)
manifest=[{'path':str(p.relative_to(src)).replace('\\','/'),'size':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in files]
(out/'文件清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
zpath=out/'puream-drama-production-package-20260915.zip'
with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
 for p in files:z.write(p,'puream-drama-production-package/'+p.relative_to(src).as_posix())
 for p in [word,out/'使用说明.md',out/'文件清单.json']:z.write(p,p.name)
with zipfile.ZipFile(zpath) as z:
 assert z.testzip() is None
 for row in manifest:assert hashlib.sha256(z.read('puream-drama-production-package/'+row['path'])).hexdigest()==row['sha256']
assert len(Document(word).paragraphs)>15
result={'skillFiles':len(files),'zipBytes':zpath.stat().st_size,'sha256':hashlib.sha256(zpath.read_bytes()).hexdigest(),'allFilesMatchSource':True,'wordReadable':True}
(out/'打包校验.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result))
