from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path(__file__).parent / '纯梦短剧老虎机-v0.16.352-简明使用说明.docx'
d = Document()
s = d.sections[0]
s.top_margin = s.bottom_margin = Cm(1.8)
s.left_margin = s.right_margin = Cm(2)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    st = d.styles[name]
    st.font.name = '微软雅黑'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
d.styles['Normal'].font.size = Pt(10.5)
d.styles['Normal'].paragraph_format.space_after = Pt(6)
d.styles['Normal'].paragraph_format.line_spacing = 1.15
d.styles['Title'].font.size = Pt(23)
d.styles['Heading 1'].font.size = Pt(15)
d.styles['Heading 1'].font.color.rgb = RGBColor.from_string('245B78')

def p(text): d.add_paragraph(text)
def h(text): d.add_heading(text, 1)
def steps(items):
    for i, text in enumerate(items, 1): p(f'{i}. {text}')

d.add_heading('纯梦短剧老虎机', 0)
p('简明使用说明  |  v0.16.352  |  2026年9月15日')
p('基本流程：新建项目 → 准备剧情与商品 → 写剧本 → 拆镜与编写提示词 → 用户确认全部提示词 → 生成资产 → 生成视频 → 剪辑导出。')
h('一、首次设置')
steps([
'打开软件，按界面提示完成授权。在左侧进入“系统设置”，填写所需的纯梦授权码，并确认相关服务余额可用。',
'在本地 AI 工作软件设置中，选择已安装、已登录的 Agent，例如 WorkBuddy。按实际需要核对各阶段的来源，刷新模型列表并选择当前账号可用的模型，不要凭显示名称猜填模型 ID。',
'检查执行入口及接入状态。若选用外部 Agent 通过 MCP 接管，需按设置页说明完成外部配置；仅打开 Agent 窗口不等于已接管。',
'配置图片来源及视频服务，核对项目与素材保存位置，点击“保存全部设置”。文本 Agent 的额度与图片、视频服务余额应分别检查。'])
h('二、选择制作模式')
t=d.add_table(rows=1, cols=2); t.style='Light Shading Accent 1'
t.rows[0].cells[0].text='模式'; t.rows[0].cells[1].text='用途与流程'
for a,b in [
('资产导入','已有完整制作包时使用，沿用包内资产和提示词。不是单纯上传一份剧本文档。'),
('资产包直投','生成或上传人物、场景、道具及商品参考资产，确认后直接生成分镜视频，跳过分镜图。'),
('首尾帧','先准备分镜的首帧、尾帧图片，再用于视频生成。'),
('分镜合图','先准备分镜合图，再按该模式生成分镜视频。')]:
    c=t.add_row().cells;c[0].text=a;c[1].text=b
p('首次使用可从“资产包直投”开始。无论选择哪种模式，后续全部生成提示词都需要用户确认。')
h('三、新建项目与准备资料')
steps([
'点击顶部“＋”新建项目，填写便于识别的项目名并选择制作模式。操作前先确认顶部当前项目正确。',
'在“剧本与商品”填写剧情需求。带货项目上传真实商品原图，填写商品名称、卖点、价格、活动和购买入口；没有活动就写“无活动”。',
'卖点留空时，AI 可根据商品名推断通常用途。具体价格、优惠和功效应以实际资料为准。故事先围绕人物和冲突展开，商品自然参与剧情。'])

d.add_page_break()
h('四、三种剧本入口')
steps([
'原创：输入剧情方向，生成选题并选择标题，再开始写剧本；也可使用“一键全流程”。',
'上传：通过上传原稿入口导入已有剧本，核对文本完整后，继续标准化与拆镜。',
'改写：打开仿写剧本窗口，上传或粘贴完整原稿；可填写人物、职业、场景等替换要求，然后开始仿写。先查看结果，再按界面操作保存。'])
p('写作期间可查看已返回正文。正文看起来写完，不代表后续拆镜、资产提取和提示词编写都已完成；请结合右侧“软件当前步骤”、本次任务状态及已保存数量判断进度。')
h('五、查看并确认全部提示词')
steps([
'剧本阶段完成后，按当前步骤按钮继续。软件将处理分镜、资产描述及各模式需要的提示词；分步制作需要按阶段继续。',
'全部提示词准备好后，会弹出“确认后续全部生成提示词”。检查人物、场景、道具、商品，以及分镜图片和视频等适用部分。英文执行稿可结合中文译文查看。',
'重点检查：对白是否完整且不重复、说话人与听者是否正确、站位和动作是否连贯、商品原图是否引用、价格活动是否准确、带货是否融入剧情。',
'需要时选择“AI 审核校正”，查看修改前后内容及修改原因，再决定是否一键应用。应用修改不等于确认制作；最后确认当前提示词版本，才继续后续生成。'])
p('重要：一键全流程也不能跳过用户确认。关闭弹窗、AI 审核完成或应用修改，都不代表您已同意开始图片、视频生成。')
h('六、生成资产、视频与成片')
steps([
'确认提示词后，进入“角色与场景”等对应环节，生成或上传参考资产。核对人物外观、场景和商品原图，必要时对具体资产重新生成。',
'资产包直投直接进入分镜视频；首尾帧、分镜合图先完成对应图片。检查当前项目的任务列表及已完成素材，再继续视频制作。',
'视频完成后进入“智能剪辑”，按界面提供的导出入口制作成片或剪映草稿。导出后实际打开检查画面、对白和声音；任务完成不等于成片内容一定符合预期。'])
h('七、遇到等待、找项目或更新')
p('等待较久：先看当前步骤、本次调用耗时、最近输出与已保存数量。界面刷新不一定代表模型有新正文；不要连续重复点击生成，以免重复提交。需要中止时使用软件的暂停或停止按钮。')
p('反馈问题：先选中出现问题的项目，点击左侧“导出运行日志”，保存 ZIP，并附上问题截图、点击的按钮和大致发生时间。日志可能含剧本和提示词，请只发给需要排查的人。')
p('查找项目：从顶部当前项目下拉框切换。若在另一种界面模式创建，先切回相应模式查找。')
p('更新软件：任务结束并保存编辑后，点击左下角版本号检查更新，或到 https://puream.cn/drama-slot-machine 下载。升级时保留项目数据目录。')

footer=s.footer.paragraphs[0]; footer.alignment=2
footer.add_run('纯梦短剧老虎机 · 简明使用说明  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
d.save(out)
check=Document(out)
assert len(check.tables)==1 and len(check.paragraphs)>30
assert '确认后续全部生成提示词' in '\n'.join(x.text for x in check.paragraphs)
print(str(out))
print(f'校验通过：{len(check.paragraphs)} 段落，{len(check.tables)} 表格；文件 {out.stat().st_size} 字节')
